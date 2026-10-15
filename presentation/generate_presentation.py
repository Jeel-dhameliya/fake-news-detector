"""TruthLens Presentation Generator — Milestone 9C Accuracy Corrections

Builds the official 9-slide PowerPoint deck in 16:9 widescreen format,
incorporating verified evaluation metrics, accurate endpoint paths (/health and /api/health),
honest failure-case demo descriptions, and distinguished findings vs. hypotheses.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE


def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Professional Academic Dark Slate & Navy Theme
    BG_COLOR = RGBColor(15, 23, 42)        # Slate 900 (#0F172A)
    CARD_BG = RGBColor(30, 41, 59)         # Slate 800 (#1E293B)
    CARD_BORDER = RGBColor(51, 65, 85)     # Slate 700 (#334155)
    TEXT_MAIN = RGBColor(248, 250, 252)    # Slate 50 (#F8FAFC)
    TEXT_MUTED = RGBColor(148, 163, 184)   # Slate 400 (#94A3B8)
    ACCENT_CYAN = RGBColor(6, 182, 212)    # Cyan 500 (#06B6D4)
    ACCENT_SKY = RGBColor(56, 189, 248)    # Sky 400 (#38BDF8)
    ACCENT_EMERALD = RGBColor(16, 185, 129)# Emerald 500 (#10B981)
    ACCENT_AMBER = RGBColor(245, 158, 11)  # Amber 500 (#F59E0B)
    ACCENT_ROSE = RGBColor(244, 63, 94)    # Rose 500 (#F43F5E)

    def apply_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()
        return bg

    def add_header(slide, tracker, title, subtitle=None):
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.2))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = tracker.upper()
        p0.font.name = "Arial"
        p0.font.size = Pt(10)
        p0.font.bold = True
        p0.font.color.rgb = ACCENT_CYAN
        p0.space_after = Pt(2)

        p1 = tf.add_paragraph()
        p1.text = title
        p1.font.name = "Arial"
        p1.font.size = Pt(22)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_MAIN

        if subtitle:
            p2 = tf.add_paragraph()
            p2.text = subtitle
            p2.font.name = "Arial"
            p2.font.size = Pt(11)
            p2.font.color.rgb = ACCENT_SKY
            p2.space_before = Pt(2)

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
        tf = card.text_frame
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.2)
        tf.margin_bottom = Inches(0.2)
        return card

    # ==========================================
    # SLIDE 1: TITLE SLIDE
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    apply_background(s1)

    badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(3.2), Inches(0.4))
    badge.fill.solid()
    badge.fill.fore_color.rgb = RGBColor(8, 47, 73)
    badge.line.color.rgb = ACCENT_CYAN
    badge.line.width = Pt(1)
    tf_b = badge.text_frame
    tf_b.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_b = tf_b.paragraphs[0]
    p_b.text = "ACADEMIC CAPSTONE PROJECT • 2026"
    p_b.font.name = "Arial"
    p_b.font.size = Pt(10)
    p_b.font.bold = True
    p_b.font.color.rgb = ACCENT_SKY
    p_b.alignment = PP_ALIGN.CENTER

    tbox = s1.shapes.add_textbox(Inches(0.8), Inches(2.1), Inches(11.733), Inches(2.5))
    tf1 = tbox.text_frame
    tf1.word_wrap = True

    p_title = tf1.paragraphs[0]
    p_title.text = "TruthLens: AI Misinformation Classifier"
    p_title.font.name = "Arial"
    p_title.font.size = Pt(36)
    p_title.font.bold = True
    p_title.font.color.rgb = TEXT_MAIN
    p_title.space_after = Pt(8)

    p_sub = tf1.add_paragraph()
    p_sub.text = "Exploring Transformer Natural Language Processing, Algorithmic Transparency, and Empirical Benchmark Evaluation"
    p_sub.font.name = "Arial"
    p_sub.font.size = Pt(16)
    p_sub.font.color.rgb = ACCENT_SKY

    card1 = add_card(s1, Inches(0.8), Inches(4.7), Inches(3.64), Inches(2.1))
    tf_c1 = card1.text_frame
    tf_c1.word_wrap = True
    p = tf_c1.paragraphs[0]
    p.text = "Inference Architecture"
    p.font.name = "Arial"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN
    p = tf_c1.add_paragraph()
    p.text = (
        "• BERT Sequence Classifier (Hugging Face)\n"
        "• Hardware execution: PyTorch MPS (Apple Silicon GPU)\n"
        "• FastAPI Async Engine & Uvicorn Backend\n"
        "• Singleton lazy loading with CPU fallback"
    )
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_MUTED

    card2 = add_card(s1, Inches(4.84), Inches(4.7), Inches(3.64), Inches(2.1))
    tf_c2 = card2.text_frame
    tf_c2.word_wrap = True
    p = tf_c2.paragraphs[0]
    p.text = "Modern Web Dashboard"
    p.font.name = "Arial"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_EMERALD
    p = tf_c2.add_paragraph()
    p.text = (
        "• React 18, TypeScript, and Vite 5 SPA\n"
        "• Dual health routes (/health & /api/health)\n"
        "• 15-second request timeouts & resilient retry\n"
        "• Transparent Model Output Score visualizer"
    )
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_MUTED

    card3 = add_card(s1, Inches(8.88), Inches(4.7), Inches(3.64), Inches(2.1))
    tf_c3 = card3.text_frame
    tf_c3.word_wrap = True
    p = tf_c3.paragraphs[0]
    p.text = "Audited Evaluation"
    p.font.name = "Arial"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_AMBER
    p = tf_c3.add_paragraph()
    p.text = (
        "• Evaluated on UCSB LIAR benchmark (1,016 rows)\n"
        "• Measured accuracy: 49.21% vs 54.72% majority baseline\n"
        "• 500 correct predictions out of 1,016 evaluated claims\n"
        "• Explicit limitation disclosures in UI"
    )
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_MUTED

    s1.notes_slide.notes_text_frame.text = (
        "SLIDE 1 — TITLE & INTRODUCTION (Time: 0:00 - 0:30)\n\n"
        "Welcome, esteemed committee members. Today I am presenting TruthLens, an academic capstone research "
        "prototype exploring automated Natural Language Processing techniques for news classification and "
        "algorithmic trust transparency. The project integrates a fine-tuned transformer classification engine "
        "powered by FastAPI with a modern, responsive React dashboard. Beyond software implementation, "
        "our research investigates the empirical limits of linguistic classifiers when evaluated out-of-domain on "
        "benchmark datasets."
    )

    # ==========================================
    # SLIDE 2: PROBLEM STATEMENT
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    apply_background(s2)
    add_header(s2, "Problem Context & Motivation", "The Misinformation Challenge & Automated Detection Pitfalls",
               "Why naive machine-learning classification cannot replace independent factual verification")

    card_left = add_card(s2, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0))
    tf_l = card_left.text_frame
    tf_l.word_wrap = True
    p = tf_l.paragraphs[0]
    p.text = "The Scale of Digital Misinformation"
    p.font.name = "Arial"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_ROSE
    p = tf_l.add_paragraph()
    p.text = (
        "\n1. Exponential Velocity:\n"
        "False claims spread up to 6x faster than verified news across digital platforms.\n\n"
        "2. Cognitive Impact:\n"
        "Sensationalized headlines and emotionally charged phrasing exploit human cognitive biases, "
        "leading readers to accept assertions without verifying sources.\n\n"
        "3. Manual Verification Bottleneck:\n"
        "Professional fact-checking organizations (e.g., PolitiFact, Snopes) require hours or days "
        "of rigorous journalistic investigation per claim, creating an urgent demand for automated assistance."
    )
    p.font.name = "Arial"
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_MUTED

    card_right = add_card(s2, Inches(6.8), Inches(1.8), Inches(5.733), Inches(5.0))
    tf_r = card_right.text_frame
    tf_r.word_wrap = True
    p = tf_r.paragraphs[0]
    p.text = "The Critical NLP Misinterpretation Trap"
    p.font.name = "Arial"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_AMBER
    p = tf_r.add_paragraph()
    p.text = (
        "\n1. Form vs. Fact:\n"
        "Language models learn statistical correlations in syntax, vocabulary, and phrasing from training datasets. "
        "They do NOT possess real-world ground truth knowledge or access live facts.\n\n"
        "2. The 'False Sense of Security' Risk:\n"
        "Many applications label softmax probabilities as 'confidence of truth'. "
        "Users mistake mathematical certainty over trained labels for verified factual reality.\n\n"
        "3. Research Motivation for TruthLens:\n"
        "Demonstrate an end-to-end classification system that balances engineering performance with "
        "scientific honesty—pairing predictions with prominent benchmark disclosures and transparent wording."
    )
    p.font.name = "Arial"
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_MUTED

    s2.notes_slide.notes_text_frame.text = (
        "SLIDE 2 — PROBLEM STATEMENT (Time: 0:30 - 1:00)\n\n"
        "The motivation behind TruthLens centers on a major dilemma in applied AI: while misinformation spreads "
        "rapidly and automated tools are desperately sought, NLP sequence models classify writing style, "
        "not factual veracity. When commercial tools present uncalibrated softmax scores as 'confidence of truth', "
        "they inadvertently generate a false sense of security. Our project addresses this problem by designing "
        "an architecture that delivers automated inference while deliberately confronting users with empirical "
        "benchmark realities and limitations."
    )

    # ==========================================
    # SLIDE 3: PROJECT OBJECTIVES
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    apply_background(s3)
    add_header(s3, "Project Goals & Scope", "Core Architectural & Research Objectives",
               "Four pillars guiding the implementation and evaluation of TruthLens")

    pillars = [
        ("1. End-to-End Pipeline", ACCENT_CYAN,
         "• Build a modular inference service using fine-tuned BERT.\n"
         "• Enable PyTorch Apple Silicon GPU acceleration (MPS) with reliable CPU fallback.\n"
         "• Implement singleton caching to eliminate redundant model loading overhead."),
        ("2. Interactive Dashboard", ACCENT_EMERALD,
         "• Develop a modern, responsive React 18 single-page application.\n"
         "• Provide real-time character counting, presets, and clean error alerts.\n"
         "• Enforce strict 15s request timeouts and 30s background health polling."),
        ("3. Algorithmic Transparency", ACCENT_AMBER,
         "• Replace ambiguous 'Confidence' with 'Model Output Score'.\n"
         "• Display raw softmax distributions without claiming calibrated factuality.\n"
         "• Prominently disclose empirical limitations directly within the UI."),
        ("4. Rigorous Evaluation", ACCENT_SKY,
         "• Benchmark the frozen model against the authoritative UCSB LIAR test split.\n"
         "• Dual health routes (/health & /api/health) reporting API liveness and model state.\n"
         "• Implement automated GitHub Actions CI and 100% offline-runnable test suites.")
    ]

    for i, (title, color, desc) in enumerate(pillars):
        col = i % 2
        row = i // 2
        x = Inches(0.8 + col * 6.0)
        y = Inches(1.8 + row * 2.6)
        c = add_card(s3, x, y, Inches(5.7), Inches(2.3))
        tf = c.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Arial"
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = color
        p = tf.add_paragraph()
        p.text = desc
        p.font.name = "Arial"
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_MUTED

    s3.notes_slide.notes_text_frame.text = (
        "SLIDE 3 — PROJECT OBJECTIVES (Time: 1:00 - 1:30)\n\n"
        "TruthLens was built upon four pillars: First, building an efficient inference engine using Hugging Face "
        "Transformers and PyTorch with hardware acceleration. Second, delivering an accessible, resilient user "
        "dashboard in React with 15-second timeouts and dual health monitoring via /health and /api/health. Third, "
        "establishing algorithmic transparency by redesigning how scores and caveats are communicated. And fourth, "
        "conducting reproducible, independent benchmark evaluations rather than relying on unverified claims."
    )

    # ==========================================
    # SLIDE 4: SYSTEM ARCHITECTURE
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    apply_background(s4)
    add_header(s4, "Engineering Design", "Decoupled Client-Server System Architecture",
               "Seamless integration of FastAPI, PyTorch MPS hardware acceleration, and React 18")

    l1 = add_card(s4, Inches(0.8), Inches(1.8), Inches(3.64), Inches(5.0))
    tf1 = l1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "Client Presentation Layer"
    p.font.name = "Arial"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN
    p = tf1.add_paragraph()
    p.text = (
        "Technology: React 18 + TypeScript + Vite 5\n\n"
        "• Component Architecture:\n"
        "  - Navbar: Health polling & model indicator\n"
        "  - InputSection: Live validation & presets\n"
        "  - ResultsSection: Output visualizer\n"
        "  - ErrorSection: HTTP 408/422/503 alerts\n\n"
        "• Resilience & Accessibility:\n"
        "  - 15s AbortController request timeout\n"
        "  - 30s background health polling with cleanup\n"
        "  - ARIA live announcements for screen readers\n"
        "  - Fully responsive from 390px mobile to 4K"
    )
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_MUTED

    l2 = add_card(s4, Inches(4.84), Inches(1.8), Inches(3.64), Inches(5.0))
    tf2 = l2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "FastAPI Backend Layer"
    p.font.name = "Arial"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_EMERALD
    p = tf2.add_paragraph()
    p.text = (
        "Technology: FastAPI + Pydantic v2 + Uvicorn\n\n"
        "• Route Endpoints & Parity:\n"
        "  - POST /api/predict (alias /predict)\n"
        "  - GET /health & GET /api/health\n\n"
        "• Strict Input Validation:\n"
        "  - 1 to 50,000 character length bounds\n"
        "  - Rejects whitespace-only or solely invisible / "
        "zero-width formatting characters\n"
        "  - HTTP 422 structured error messages\n\n"
        "• Health & Liveness Separation:\n"
        "  - Reports api_liveness without loading model\n"
        "  - Reports model_readiness ('unloaded'/'ready')"
    )
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_MUTED

    l3 = add_card(s4, Inches(8.88), Inches(1.8), Inches(3.64), Inches(5.0))
    tf3 = l3.text_frame
    tf3.word_wrap = True
    p = tf3.paragraphs[0]
    p.text = "Inference & Deep Learning"
    p.font.name = "Arial"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_AMBER
    p = tf3.add_paragraph()
    p.text = (
        "Technology: PyTorch + Hugging Face Transformers\n\n"
        "• Model Specification:\n"
        "  - ID: dima806/fake-news-classifier\n"
        "  - Base: BertForSequenceClassification\n"
        "  - Tokenizer: BERT Cased (max 512 tokens)\n\n"
        "• Hardware Execution:\n"
        "  - Hardware acceleration via MPS (Apple Silicon)\n"
        "  - Automatic fallback to CPU on exceptions\n\n"
        "• Singleton Lifecycle:\n"
        "  - Lazy loading on first inference request\n"
        "  - Zero re-initialization overhead on repeats"
    )
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_MUTED

    s4.notes_slide.notes_text_frame.text = (
        "SLIDE 4 — SYSTEM ARCHITECTURE (Time: 1:30 - 2:00)\n\n"
        "The architecture is structured across three decoupled layers: At the presentation tier, a Vite-powered "
        "React SPA manages user interactions and enforces client-side timeouts. At the middle tier, FastAPI provides "
        "asynchronous request handling, schema validation via Pydantic v2, and health check route parity across "
        "/health and /api/health. At the machine learning tier, our FakeNewsClassifier service encapsulates a fine-tuned "
        "BERT transformer managed as an in-memory singleton. It leverages Apple Silicon Metal Performance Shaders for "
        "hardware acceleration with dynamic CPU fallback."
    )

    # ==========================================
    # SLIDE 5: HOW IT WORKS
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    apply_background(s5)
    add_header(s5, "Inference Pipeline & Mechanics", "From Raw Text to Prediction Verdict",
               "Step-by-step trace of how TruthLens processes a submitted claim")

    steps = [
        ("Step 1: Input Validation", ACCENT_CYAN,
         "• Request received at POST /api/predict (alias /predict)\n"
         "• Checks length (1 - 50,000 chars)\n"
         "• Rejects whitespace-only or solely invisible / zero-width characters (e.g. U+200B, U+FEFF)\n"
         "• Rejects empty or invalid payloads with HTTP 422"),
        ("Step 2: Tokenization", ACCENT_SKY,
         "• BERT WordPiece tokenizer decomposes text\n"
         "• Appends special tokens [CLS] and [SEP]\n"
         "• Truncates sequences exceeding 512 tokens\n"
         "• Generates input_ids and attention_mask tensors"),
        ("Step 3: Forward Pass & Logits", ACCENT_AMBER,
         "• Tensors dispatched to active device (MPS or CPU)\n"
         "• 12 transformer encoder layers compute contextual embeddings\n"
         "• Linear classification head projects [CLS] pooler output\n"
         "• Generates 2 unnormalized raw logits: [z_fake, z_real]"),
        ("Step 4: Softmax & Verdict", ACCENT_EMERALD,
         "• Softmax normalizes logits: p_i = exp(z_i) / Σ exp(z_j)\n"
         "• Dynamically resolves class mapping (index 0=FAKE, 1=REAL)\n"
         "• Returns JSON payload with predicted label and class probabilities\n"
         "• UI renders verdict card and benchmark limitation notices")
    ]

    for i, (stitle, scolor, sdesc) in enumerate(steps):
        x = Inches(0.8 + i * 2.96)
        c = add_card(s5, x, Inches(1.8), Inches(2.85), Inches(4.0))
        tf = c.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = stitle
        p.font.name = "Arial"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = scolor
        p = tf.add_paragraph()
        p.text = sdesc
        p.font.name = "Arial"
        p.font.size = Pt(10.5)
        p.font.color.rgb = TEXT_MUTED

    bot_card = add_card(s5, Inches(0.8), Inches(6.0), Inches(11.733), Inches(0.9), bg_color=RGBColor(26, 36, 61))
    tf_b = bot_card.text_frame
    tf_b.word_wrap = True
    p = tf_b.paragraphs[0]
    p.text = "CRUCIAL DISTINCTION: CLASSIFICATION VS. FACT-CHECKING"
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = ACCENT_AMBER
    p = tf_b.add_paragraph()
    p.text = "The pipeline evaluates whether text exhibits linguistic patterns correlating with fine-tuned training labels. It does NOT search the web, query knowledge bases, or verify external evidence."
    p.font.name = "Arial"
    p.font.size = Pt(10.5)
    p.font.color.rgb = TEXT_MAIN

    s5.notes_slide.notes_text_frame.text = (
        "SLIDE 5 — HOW IT WORKS (Time: 2:00 - 2:30)\n\n"
        "Here is the execution trace: Step 1 validates the payload against length limits (1 to 50,000 characters) and "
        "rejects inputs that are empty, whitespace-only, or composed entirely of invisible or zero-width formatting characters. "
        "Step 2 tokenizes the text into WordPiece tokens with a 512-token ceiling. In Step 3, the tensors pass through "
        "BERT's 12 self-attention layers to produce unnormalized logits on MPS or CPU. Finally, Step 4 applies softmax to derive output "
        "scores over FAKE and REAL classes. We emphasize the bottom callout: this is mathematical text classification, "
        "not evidentiary fact-checking."
    )

    # ==========================================
    # SLIDE 6: LIVE DEMO & DEMONSTRATION SAMPLES
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    apply_background(s6)
    add_header(s6, "Live Application Demonstration", "Demonstrating the Model's Behavior on Diverse Samples",
               "Why high confidence on false claims illustrates an architectural limitation rather than factual proof")

    card_s1 = add_card(s6, Inches(0.8), Inches(1.8), Inches(5.7), Inches(4.3))
    tf_s1 = card_s1.text_frame
    tf_s1.word_wrap = True
    p = tf_s1.paragraphs[0]
    p.text = "Demo Sample 1: Scientific Excerpt"
    p.font.name = "Arial"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_EMERALD
    p = tf_s1.add_paragraph()
    p.text = (
        "Input Statement:\n"
        "\"NASA confirmed today that the James Webb Space Telescope detected water vapor in the atmosphere of a rocky exoplanet orbiting within its host star habitable zone.\"\n\n"
        "Observed Live Output:\n"
        "• Predicted Label: REAL\n"
        "• Model Output Score: 100.0% (REAL: 100.0%, FAKE: 0.0%)\n"
        "• Inference Execution: PyTorch on Apple Silicon MPS\n\n"
        "Observed Result:\n"
        "The model classified this structured scientific statement as REAL with 100.0% model output score."
    )
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_MUTED

    card_s2 = add_card(s6, Inches(6.8), Inches(1.8), Inches(5.733), Inches(4.3))
    tf_s2 = card_s2.text_frame
    tf_s2.word_wrap = True
    p = tf_s2.paragraphs[0]
    p.text = "Demo Sample 2: Conspiracy Headline"
    p.font.name = "Arial"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_ROSE
    p = tf_s2.add_paragraph()
    p.text = (
        "Input Statement:\n"
        "\"SHOCKING PROOF: Government secretly installing microchips in tap water to control thoughts of citizens across nationwide network! Whistleblowers reveal explosive hidden documents!\"\n\n"
        "Observed Live Output:\n"
        "• Predicted Label: REAL\n"
        "• Model Output Score: 100.0% (REAL: 100.0%, FAKE: 0.0%)\n"
        "• Inference Execution: PyTorch on Apple Silicon MPS\n\n"
        "Demonstrated Classifier Failure Case:\n"
        "Despite making an unfounded conspiracy claim, the model returned REAL with 100.0% score. This observed run demonstrates a failure case for the current classifier: isolated statistical text classifiers without external evidence retrieval or verification mechanisms cannot determine factual reality."
    )
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_MUTED

    bot_d = add_card(s6, Inches(0.8), Inches(6.25), Inches(11.733), Inches(0.75), bg_color=RGBColor(26, 36, 61))
    tf_bd = bot_d.text_frame
    tf_bd.word_wrap = True
    p = tf_bd.paragraphs[0]
    p.text = "UI Mitigation in TruthLens: The dashboard prominently warns the user that 'Model Output Score' reflects statistical label preference, and displays our LIAR benchmark disclosure directly beneath the verdict."
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.color.rgb = ACCENT_SKY

    s6.notes_slide.notes_text_frame.text = (
        "SLIDE 6 — LIVE DEMO & DEMONSTRATED FAILURE CASE (Time: 2:30 - 3:15)\n\n"
        "During live testing in the recorded evaluation run, we tested two contrasting samples: Sample 1 is a scientific report "
        "about NASA's Webb telescope, which the model classified as REAL with a 100.0% Model Output Score. However, Sample 2 "
        "is a conspiracy headline asserting that the government installs microchips in tap water to control citizens' thoughts. "
        "In the observed run, the model ALSO classified this statement as REAL with a 100.0% output score. This demonstrates a "
        "critical failure case for the current classifier: because it operates solely on statistical text patterns without external "
        "evidence retrieval or fact-checking knowledge bases, high output scores cannot be treated as evidence of factual truth. "
        "This empirical outcome directly justifies why TruthLens labels scores as 'Model Output Score' rather than factual truth."
    )

    # ==========================================
    # SLIDE 7: BENCHMARK METHODOLOGY
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    apply_background(s7)
    add_header(s7, "Rigorous Out-of-Domain Evaluation", "LIAR Benchmark Methodology & Confusion Matrix",
               "Evaluating transfer capability on 1,016 political claims from William Yang Wang (ACL 2017)")

    c_proto = add_card(s7, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0))
    tf_p = c_proto.text_frame
    tf_p.word_wrap = True
    p = tf_p.paragraphs[0]
    p.text = "Benchmark Protocol & Accounting"
    p.font.name = "Arial"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN
    p = tf_p.add_paragraph()
    p.text = (
        "• Authoritative Source:\n"
        "  UCSB LIAR Original Publication Archive (ACL 2017)\n"
        "  Verified SHA-256 Checksum:\n"
        "  611c1addad919743dde15822b87a60bfb760d8f85597f25289e34621800654c7\n\n"
        "• Row Accounting (test.tsv):\n"
        "  - Total Test Rows: 1,283\n"
        "  - Excluded 'half-true' Rows: 267 (ambiguous)\n"
        "  - Evaluated Rows: 1,016 (pure text statement only)\n\n"
        "• Operational Binary Mapping:\n"
        "  - FAKE: false (250) + pants-fire (92) + barely-true (214) = 556 (54.72%)\n"
        "  - REAL: true (211) + mostly-true (249) = 460 (45.28%)\n\n"
        "• Reproducibility:\n"
        "  Automated script backend/scripts/evaluate_model.py\n"
        "  Outputs compact prediction pairings for external auditing."
    )
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_MUTED

    c_cm = add_card(s7, Inches(6.8), Inches(1.8), Inches(5.733), Inches(5.0))
    tf_cm = c_cm.text_frame
    tf_cm.word_wrap = True
    p = tf_cm.paragraphs[0]
    p.text = "Measured Confusion Matrix (N = 1,016)"
    p.font.name = "Arial"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_SKY

    p = tf_cm.add_paragraph()
    p.text = (
        "\nMatrix Layout (Rows: Actual Class | Columns: Predicted Class):\n\n"
        "                     Predicted FAKE      Predicted REAL      Total Actual\n"
        "Actual FAKE:              416 (TP)            140 (FN)           556\n"
        "Actual REAL:              376 (FP)             84 (TN)           460\n"
        "Total Predicted:          792                 224              1,016\n\n"
        "Correct Predictions: 416 (TP) + 84 (TN) = 500 / 1,016 (49.21%)\n\n"
        "Critical Empirical Takeaways:\n"
        "1. Extreme False Positive Rate:\n"
        "   376 out of 460 actual REAL claims were falsely classified as FAKE.\n"
        "2. Skewed Prediction Distribution:\n"
        "   The model predicted FAKE on 792 out of 1,016 claims (77.95%).\n"
        "3. Low REAL Recall:\n"
        "   The model correctly identified only 84 REAL claims (18.26% recall)."
    )
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_MUTED

    s7.notes_slide.notes_text_frame.text = (
        "SLIDE 7 — BENCHMARK METHODOLOGY & CONFUSION MATRIX (Time: 3:15 - 3:45)\n\n"
        "To test our model objectively, we benchmarked it against the LIAR dataset using the official UCSB archive "
        "verified by SHA-256 checksum. Out of 1,283 test items, 267 ambiguous half-true statements were excluded following "
        "the documented binary evaluation protocol, leaving 1,016 evaluated claims. Looking at the confusion matrix, "
        "there were 416 True Positives for FAKE and 84 True Negatives for REAL, totaling 500 correct classifications out of "
        "1,016 (49.21%). The model predicted FAKE 792 times (77.95%), misclassifying 376 actual REAL statements as FAKE "
        "and yielding a low REAL recall of 18.26%."
    )

    # ==========================================
    # SLIDE 8: RESULTS & LIMITATIONS
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    apply_background(s8)
    add_header(s8, "Measured Metrics & Academic Honesty", "Benchmark Results vs. Majority Baseline",
               "Documenting the gap between headline claims and out-of-domain reality")

    c_met = add_card(s8, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0))
    tf_m = c_met.text_frame
    tf_m.word_wrap = True
    p = tf_m.paragraphs[0]
    p.text = "Empirical Evaluation Metrics"
    p.font.name = "Arial"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_AMBER

    p = tf_m.add_paragraph()
    p.text = (
        "\nMetric                     Score       Majority Baseline Comparison\n"
        "----------------------------------------------------------\n"
        "Overall Accuracy           49.21%      BELOW Baseline (54.72%)\n"
        "                                       (500 / 1,016 correct predictions)\n"
        "Majority Baseline          54.72%      (Naive 'Always FAKE' heuristic)\n"
        "Macro F1                   0.4314      Unweighted mean of class F1s\n"
        "FAKE Precision / Recall    0.5253 / 0.7482\n"
        "REAL Precision / Recall    0.3750 / 0.1826 (Identifies 84 of 460)\n\n"
        "Controlled Research Audits (Milestone 6):\n"
        "• Threshold Tuning: Shifting boundary from 0.50 to 0.13 marginally raised "
        "accuracy to 52.85%, still below the 54.72% majority baseline.\n"
        "• Candidate Screening: Tested RoBERTa and LIAR-specialized BERT on validation; "
        "neither exceeded validation majority baseline (59.46%)."
    )
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_MUTED

    c_lim = add_card(s8, Inches(6.8), Inches(1.8), Inches(5.733), Inches(5.0))
    tf_lm = c_lim.text_frame
    tf_lm.word_wrap = True
    p = tf_lm.paragraphs[0]
    p.text = "Findings vs. Hypothesized Causes"
    p.font.name = "Arial"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_ROSE

    p = tf_lm.add_paragraph()
    p.text = (
        "\n1. Measured Outcome Below Baseline (Verified Finding):\n"
        "The model's 49.21% accuracy falls below the 54.72% majority-class baseline on the evaluated 1,016-claim LIAR test subset.\n\n"
        "2. Absence of External Evidence (Architectural Limitation):\n"
        "Isolated language models evaluate lexical and syntactic patterns without external retrieval; factual assertions cannot be verified from phrasing alone.\n\n"
        "3. Uncalibrated Model Output Scores (Measured Limitation):\n"
        "Softmax output scores frequently saturate near 1.0 or 0.0, which cannot be interpreted as calibrated probabilities of factual truth.\n\n"
        "4. Domain Shift & Training Provenance (Leading Hypothesis):\n"
        "Because the model author did not document training corpora or preprocessing scripts, domain divergence on short political statements is a hypothesized factor, not a proven cause."
    )
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_MUTED

    s8.notes_slide.notes_text_frame.text = (
        "SLIDE 8 — RESULTS & LIMITATIONS (Time: 3:45 - 4:20)\n\n"
        "The empirical findings are definitive: the model achieved 49.21% overall accuracy, with 500 correct predictions "
        "(416 + 84) out of 1,016 evaluated samples. This is 5.51 percentage points below the 54.72% majority-class baseline—meaning "
        "a naive heuristic that simply guessed 'FAKE' on every claim would outperform the neural network on this dataset. "
        "Furthermore, our research audits in Milestone 6 confirmed that threshold tuning only reached 52.85%, still below baseline. "
        "We distinguish measured facts from hypotheses: the absence of external evidence retrieval and score uncalibration are "
        "demonstrated structural limitations, while domain shift from undocumented training corpora remains a leading hypothesis. "
        "Presenting these numbers honestly is a core contribution of our capstone."
    )

    # ==========================================
    # SLIDE 9: CONCLUSION & FUTURE WORK
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    apply_background(s9)
    add_header(s9, "Summary & Future Roadmap", "Project Contributions & Research Directions",
               "What TruthLens accomplished and the technical roadmap for real-world reliability")

    c_built = add_card(s9, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0))
    tf_b = c_built.text_frame
    tf_b.word_wrap = True
    p = tf_b.paragraphs[0]
    p.text = "What TruthLens Successfully Built"
    p.font.name = "Arial"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_EMERALD
    p = tf_b.add_paragraph()
    p.text = (
        "\n✓ Complete, decoupled full-stack prototype:\n"
        "  FastAPI backend with PyTorch MPS GPU inference and React 18 responsive dashboard.\n\n"
        "✓ Production reliability engineering:\n"
        "  Lazy singleton loading, 15s request timeouts, 30s background health polling, "
        "  and robust invisible character input validation.\n\n"
        "✓ 100% automated verification & CI:\n"
        "  73 backend tests (offline-verified), 12 frontend tests, and automated GitHub Actions CI.\n\n"
        "✓ Trust transparency pioneer:\n"
        "  First-class UI disclosures separating model preference scores from factual truth."
    )
    p.font.name = "Arial"
    p.font.size = Pt(11.5)
    p.font.color.rgb = TEXT_MUTED

    c_fut = add_card(s9, Inches(6.8), Inches(1.8), Inches(5.733), Inches(5.0))
    tf_f = c_fut.text_frame
    tf_f.word_wrap = True
    p = tf_f.paragraphs[0]
    p.text = "Roadmap for Future Research"
    p.font.name = "Arial"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN
    p = tf_f.add_paragraph()
    p.text = (
        "\n1. Retrieval-Augmented Generation (RAG):\n"
        "Integrate search APIs and ClaimReview fact-check databases to retrieve verified context "
        "before classifying claims.\n\n"
        "2. Temperature Scaling & Calibration:\n"
        "Apply Platt scaling or temperature scaling on validation logits to align output scores "
        "with actual empirical accuracy.\n\n"
        "3. Explainable AI (XAI) Attribution:\n"
        "Implement LIME or SHAP token saliency heatmaps to show users which words influenced the decision.\n\n"
        "4. Source Reputation Heuristics:\n"
        "Incorporate publisher credibility signals and domain reputation metadata."
    )
    p.font.name = "Arial"
    p.font.size = Pt(11.5)
    p.font.color.rgb = TEXT_MUTED

    s9.notes_slide.notes_text_frame.text = (
        "SLIDE 9 — CONCLUSION & ROADMAP (Time: 4:20 - 5:00)\n\n"
        "In conclusion, TruthLens successfully delivered an end-to-end AI software prototype with production-grade "
        "engineering—including dual health check parity across /health and /api/health, Apple Silicon GPU acceleration, "
        "15-second request timeouts, and 100% passing automated CI. More importantly, the project bridges the gap between software "
        "implementation and scientific responsibility by demonstrating why standalone linguistic classifiers cannot autonomously "
        "determine factual truth. Our proposed roadmap outlines the transition to Retrieval-Augmented Generation, probability "
        "calibration, and token-level explainability. Thank you for your time, and I look forward to answering your questions."
    )

    output_path = "presentation/truthlens_presentation.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    create_deck()
